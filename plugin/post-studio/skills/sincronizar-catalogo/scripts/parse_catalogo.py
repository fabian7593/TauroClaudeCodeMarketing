# -*- coding: utf-8 -*-
"""
Genera el espejo local JSON del catálogo maestro a partir de la hoja fuente.

Dos vías, según lo que haya disponible:

  1) --xlsx <archivo.xlsx>   (RECOMENDADA)
     Lee el Excel directo con openpyxl. No pasa una sola celda por el contexto
     del modelo, mapea por NOMBRE de columna (aguanta que agreguen/muevan
     columnas) y es lo más barato y confiable.

  2) --dump <volcado.txt>    (RESPALDO)
     Usa el volcado de texto que deja el conector de Drive al leer el archivo,
     cuando no hay copia local del Excel. Mapea por POSICIÓN de columna, así que
     si el Excel cambia de estructura hay que actualizar `columnas` en el config.

Uso:
    python parse_catalogo.py --config <catalogo.config.json> --xlsx <archivo.xlsx>
    python parse_catalogo.py --config <catalogo.config.json> --dump <volcado.txt>
"""
import argparse
import csv
import datetime
import io
import json
import os
import re
import sys
import unicodedata


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------
def normalizar(texto):
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return " ".join(texto.lower().split())


def cargar_config(ruta):
    with io.open(ruta, encoding="utf-8") as fh:
        return json.load(fh)


def a_bool(valor):
    return normalizar(valor) in ("si", "s", "yes", "true", "x", "1")


def a_numero(valor, tipo):
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return int(valor) if tipo == "int" else float(valor)
    valor = str(valor).strip().replace("%", "")
    if not valor:
        return None
    try:
        return int(float(valor.replace(",", "."))) if tipo == "int" else float(valor.replace(",", "."))
    except ValueError:
        return None


def limpiar(fila, cfg):
    """Aplica tipos, deriva campos y saca las columnas que no van al espejo."""
    for nombre in cfg.get("booleanos", []):
        fila[nombre] = a_bool(fila.get(nombre))
    for nombre, tipo in cfg.get("numericos", {}).items():
        fila[nombre] = a_numero(fila.get(nombre), tipo)
    for nombre, sep in cfg.get("listas", {}).items():
        crudo = fila.get(nombre) or ""
        fila[nombre] = [p.strip() for p in str(crudo).split(sep) if p.strip()]

    fila["esAnime"] = fila.get("categoria") in cfg.get("categoriasAnime", [])
    for nombre in cfg.get("columnasOmitidas", []):
        fila.pop(nombre, None)

    tipados = set(cfg.get("booleanos", [])) | set(cfg.get("numericos", {})) | set(cfg.get("listas", {}))
    for nombre in list(fila):
        valor = fila[nombre]
        if isinstance(valor, str):
            valor = valor.strip()
            fila[nombre] = valor or None
        elif isinstance(valor, (int, float)) and nombre not in tipados and not isinstance(valor, bool):
            # el Excel devuelve numéricas las celdas que en realidad son códigos
            # (TMDB ID, etc.): sin esto quedarían como "890.0" y romperían las URLs.
            fila[nombre] = str(int(valor)) if float(valor).is_integer() else str(valor)
    return fila


# --------------------------------------------------------------------------
# vía 1: Excel local (recomendada)
# --------------------------------------------------------------------------
def leer_xlsx(ruta, cfg, avisos):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("Falta openpyxl. Instalalo con: pip install openpyxl")

    fuente = cfg["fuente"]
    wb = load_workbook(ruta, read_only=True, data_only=True)
    nombre_hoja = fuente.get("hoja")
    if nombre_hoja not in wb.sheetnames:
        objetivo = normalizar(nombre_hoja)
        candidatas = [h for h in wb.sheetnames if normalizar(h) == objetivo]
        if not candidatas:
            raise SystemExit("La hoja %r no está en el Excel. Hojas disponibles: %s"
                             % (nombre_hoja, ", ".join(wb.sheetnames)))
        nombre_hoja = candidatas[0]
    hoja = wb[nombre_hoja]

    mapa = {normalizar(k): v for k, v in cfg.get("mapaColumnas", {}).items()}
    if not mapa:
        raise SystemExit("El config no tiene 'mapaColumnas' (necesario para leer el Excel).")

    columnas = None
    filas = []
    for celdas in hoja.iter_rows(values_only=True):
        valores = ["" if c is None else c for c in celdas]
        if columnas is None:
            # la fila de encabezado es la primera que mapea al menos 5 columnas conocidas
            posibles = [mapa.get(normalizar(v)) for v in valores]
            if sum(1 for p in posibles if p) >= 5:
                columnas = posibles
                desconocidas = [str(v) for v, p in zip(valores, posibles) if v and not p]
                if desconocidas:
                    avisos.append("columnas del Excel ignoradas (no están en mapaColumnas): %s"
                                  % ", ".join(desconocidas))
            continue

        fila = {}
        for campo, valor in zip(columnas, valores):
            if campo:
                fila[campo] = valor if not isinstance(valor, str) else valor.strip()
        if not fila.get("vtxId"):
            continue
        filas.append(limpiar(fila, cfg))

    wb.close()
    if columnas is None:
        raise SystemExit("No se encontró la fila de encabezados en la hoja %r." % nombre_hoja)
    return filas


# --------------------------------------------------------------------------
# vía 2: volcado de texto del conector de Drive (respaldo)
# --------------------------------------------------------------------------
def leer_dump(ruta, cfg, avisos):
    with io.open(ruta, encoding="utf-8") as fh:
        texto = fh.read()
    try:
        contenido = json.loads(texto)["fileContent"]
    except (ValueError, KeyError, TypeError):
        contenido = texto

    fuente = cfg["fuente"]
    ancla = fuente["anclaHoja"]
    if ancla not in contenido:
        raise SystemExit("No se encontró el ancla de la hoja (%r) en el volcado." % ancla)
    hoja = contenido[contenido.index(ancla):]
    patron = "(?= %s[0-9]{%d},)" % (re.escape(fuente["prefijoId"]), fuente["digitosId"])
    bloques = [b.strip() for b in re.split(patron, hoja)[1:] if b.strip()]

    columnas = cfg["columnas"]
    filas = []
    for bloque in bloques:
        valores = next(csv.reader([bloque]))
        if len(valores) < len(columnas):
            avisos.append("fila corta (%d de %d campos): %s"
                          % (len(valores), len(columnas), valores[0]))
            valores += [""] * (len(columnas) - len(valores))
        elif len(valores) > len(columnas):
            cola = ",".join(valores[len(columnas) - 1:])
            valores = valores[:len(columnas) - 1] + [cola]
        fila = {c: (v.strip() if isinstance(v, str) else v) for c, v in zip(columnas, valores)}
        filas.append(limpiar(fila, cfg))
    return filas


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", required=True)
    ap.add_argument("--xlsx", help="Excel local (vía recomendada)")
    ap.add_argument("--dump", help="volcado de texto del conector de Drive (respaldo)")
    ap.add_argument("--overrides", help="correcciones locales al catálogo (default: catalogo.overrides.json)")
    ap.add_argument("--salida")
    args = ap.parse_args()

    cfg = cargar_config(args.config)
    avisos = []

    origen = args.xlsx or cfg["fuente"].get("archivoLocal")
    if origen and os.path.isfile(origen):
        titulos = leer_xlsx(origen, cfg, avisos)
        via = "xlsx local (%s)" % origen
    elif args.dump:
        titulos = leer_dump(args.dump, cfg, avisos)
        via = "volcado de Drive (%s)" % args.dump
    else:
        raise SystemExit(
            "No hay de dónde leer. Pasá --xlsx <archivo> (recomendado) o --dump <volcado.txt>.\n"
            "Si el Excel debería estar local, revisá 'fuente.archivoLocal' en el config.")

    # correcciones locales: colecciones que el Excel todavía no tiene cargadas.
    # Se aplican DESPUÉS de leer la hoja, así sobreviven a cada sincronización.
    ruta_ov = args.overrides or cfg.get("overrides", "catalogo.overrides.json")
    if os.path.isfile(ruta_ov):
        with io.open(ruta_ov, encoding="utf-8") as fh:
            overrides = json.load(fh)
        por_id = {t["vtxId"]: t for t in titulos}
        aplicados = 0
        for nombre, ids in (overrides.get("colecciones") or {}).items():
            for vtx in ids:
                t = por_id.get(vtx)
                if not t:
                    avisos.append("override: %s no existe en el catálogo (colección %s)" % (vtx, nombre))
                    continue
                if t.get("coleccion") and t["coleccion"] != nombre:
                    avisos.append("override: %s ya tenía colección %r, se reemplaza por %r"
                                  % (vtx, t["coleccion"], nombre))
                t["coleccion"] = nombre
                t["coleccionOrigen"] = "override"
                aplicados += 1
        if aplicados:
            print("overrides aplicados: %d titulos (%s)" % (aplicados, ruta_ov))

    colecciones = {}
    for t in titulos:
        if t.get("coleccion"):
            colecciones.setdefault(t["coleccion"], []).append(t["vtxId"])

    salida = args.salida or cfg.get("salida", "catalogo.json")
    datos = {
        "_comment": (
            "Espejo local del catálogo maestro, generado por el skill sincronizar-catalogo. "
            "NO se edita a mano: se regenera desde la hoja fuente cada vez que el catálogo cambia. "
            "El estado de producción (qué pieza ya se creó/publicó) va aparte, en estado-catalogo.json."
        ),
        "fuente": dict(cfg["fuente"], generado=datetime.date.today().isoformat(), via=via),
        "totalTitulos": len(titulos),
        "totalColecciones": len(colecciones),
        "colecciones": {k: colecciones[k] for k in sorted(colecciones)},
        "titulos": titulos,
    }
    if avisos:
        datos["avisos"] = avisos

    carpeta = os.path.dirname(os.path.abspath(salida))
    if carpeta and not os.path.isdir(carpeta):
        os.makedirs(carpeta)
    with io.open(salida, "w", encoding="utf-8") as fh:
        fh.write(json.dumps(datos, ensure_ascii=False, indent=2) + "\n")

    print("%s: %d titulos, %d colecciones (via %s)" % (salida, len(titulos), len(colecciones), via))
    for aviso in avisos:
        print("aviso: %s" % aviso, file=sys.stderr)


if __name__ == "__main__":
    main()
